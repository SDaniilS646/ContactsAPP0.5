# -*- coding: utf-8 -*

from openpyxl import Workbook, load_workbook
import openpyxl.styles as style
import pandas as pd

import sys
import os
from pathlib import Path

this_file_path = Path(__file__).parent
project_path = this_file_path.parent
sys.path.append(str(this_file_path.parent))

from ml_backend.work_classifier import test_model
from utils.helpers import printLabel, date_to_filename




from datetime import datetime

from collections import Counter

import re

def pyxl_color(color):
  return color.lstrip("#")

def xl_test(filepath):

  if os.path.exists(filepath):
    wb = load_workbook(filepath)
    ws = wb.active
    ws.cell(1, 1, value='old_file')
  else:
    wb = Workbook()
    ws = wb.active
    ws.cell(1, 1, value='new_file')

  font_style = style.Font(bold=True, size=14, color=pyxl_color("#FFFFFF"), name="Times New Roman")
  text_background = style.PatternFill(start_color=pyxl_color("#180052"), end_color=pyxl_color("#180052"), fill_type="solid")
  center_alignment = style.Alignment(horizontal='center', vertical='center', wrap_text=True)

  ws.cell(1, 1).font = font_style
  ws.cell(1, 1).fill = text_background
  ws.cell(1, 1).alignment = center_alignment

  wb.save(filepath)

def xl_get_last_row(sh):
  for row in range(sh.max_row, 0, -1):
    if any(sh.cell(row=row, column=col).value is not None for col in range(1, sh.max_column + 1)): return row
  
  return 0

def get_column(row, col_name, default_col=-1):
  for i, cell in enumerate(row):
    if re.match(col_name.lower().strip(), cell.lower().strip()):
      return i
  return default_col

def process_vor_file(file_data, chosen_sections=None, output_sections=None, output_categories=None, output_table_data=None, output_config=None, progress_callback=None):

  if chosen_sections is None or file_data['sections'] is None:
    return
  if progress_callback:
    from PyQt6.QtWidgets import QApplication

  temp_filepath = file_data['file_path']
  sections = file_data['sections']

  # objectName = obj_name if obj_name else 'Новый объект'

  wb = load_workbook(temp_filepath, read_only=True, data_only=True)

  ws = wb[wb.sheetnames[0]]
  xl_lastrow = xl_get_last_row(ws)

  for i, row in enumerate(ws.iter_rows(min_row=1, max_row=xl_lastrow, values_only=True)):
    if row[0] == "№ п.п.": 

      pp_col = get_column(row, '№ п.п.', 0)
      name_col = get_column(row, 'Наименование работ', 1)
      meas_col = get_column(row, 'Ед. изм.', 2)
      amount_col = get_column(row, 'Объем', 3)
      type_col = get_column(row, 'Тип позиции', 4)
  
  # for i, choice in enumerate(chosen_sections):
  #   chosen_sections[i] = int(choice)

  print(sections)
  printLabel('START PROCESSING')

  output_sect_id = 1
  output_cat_id = 1
  output_table_data_id = 1
  output_config_id = 1

  for i in chosen_sections:
    if progress_callback:
      now = datetime.now()
      progress_callback(0, f'{now} - Раздел {sections[i][0]}')
      # QApplication.processEvents()

    output_sections.append({'id':output_sect_id, 'name':sections[i][0], 'parent_id':None})
    
    
    pred_categories = []
    work_num = 1
    
    is_work = False

    for row_num, row in enumerate(ws.iter_rows(min_row=sections[i][1], max_row=sections[i+1][1] if i + 1 < len(sections) else xl_lastrow, values_only=True)):
      
      if progress_callback:
        if i + 1 < len(sections):          
          progress_perc = row_num * 100 / sections[i][1]
          progress_callback(int(progress_perc) if int(progress_perc) < 100 else 100)
          # QApplication.processEvents()
        else:
          progress_perc = (row_num + sections[i][1]) * 100 / xl_lastrow
          progress_callback(int(progress_perc) if int(progress_perc) < 100 else 100)
          # QApplication.processEvents()

      if row[type_col] is not None and row[type_col].lower().strip() == "Работа".lower().strip():
        if is_work == True:
          work_material_arr = []
          is_work = False

        material_num = 0

        work_name = row[name_col]
        pred_category = test_model([work_name], True)

        cur_cat_id = None
        
        if str(pred_category) not in pred_categories:
          output_categories.append({
            'id':output_cat_id, 
            'name': str(pred_category), 
            'section_id': output_sect_id
            })
          
          output_config.append({
            'id': output_config_id,
            'section_id': output_sect_id,
            'category_id': output_cat_id,
            'include_works': True,
            'sum_materials': False,
            'additional_files': []
          })

          cur_cat_id = output_cat_id
          output_config_id += 1
          output_cat_id += 1
        else:
          for cat in output_categories:
            if cat['name'] == pred_category:
              cur_cat_id = cat['id']
              break
        
        output_table_data.append({
          'id': output_table_data_id, 
          'category_id': cur_cat_id, 
          'data': {'code': work_num, 
                   'name': row[name_col].strip(), 
                   'chars': '', 
                   'measure': row[meas_col], 
                   'amount': row[amount_col], 
                   'type': row[type_col],
                   'include': True}
          })
        
        # output_config.append({'id': output_config_id, 'category'})
        output_table_data_id += 1

        pred_categories.append(pred_category)

        print(f"\nPrediction result: {pred_category}\n")
        work_num = Counter(pred_categories)[pred_category]
        work_material_arr = [pred_category, [work_num, row[name_col].strip(), '', row[meas_col], row[amount_col]]]
        is_work = True
      
      elif row[type_col] is not None and row[type_col].lower().strip() == "Материал".lower().strip() and is_work == True:
        material_num += 1
        output_table_data.append({
          'id': output_table_data_id, 
          'category_id': cur_cat_id, 
          'data': {'code': f"{work_num}.{material_num}", 
                   'name': row[name_col].strip(), 
                   'chars': '', 
                   'measure': row[meas_col], 
                   'amount': row[amount_col], 
                   'type': row[type_col],
                   'include': True}
          # [f"{work_num}.{material_num}", row[name_col].strip(), '', row[meas_col], row[amount_col], row[type_col]]
          })
        
        output_table_data_id += 1
        work_material_arr.append([f"{work_num}.{material_num}", row[name_col].strip(), '', row[meas_col], row[amount_col]])

      elif row[type_col] is None and is_work == True:
        work_material_arr = []
        is_work = False
        
      elif row_num + sections[i][1] == xl_lastrow and is_work == True:
        work_material_arr = []
        is_work = False

    # if status_update:
    #   status_update.setText(status_update.text() + '✅')
    #   QApplication.processEvents()
    
    output_sect_id += 1
        
  wb.close()
  os.remove(temp_filepath)
  file_data['file_path'] = None

def get_vor_file(filepath, statistics_lbl=None, file_data=None):

  # objectName = obj_name if obj_name else 'Новый объект'

  printLabel("Loading VOR")
  if statistics_lbl:
    from PyQt6.QtWidgets import QApplication

  df = pd.read_excel(filepath, sheet_name=0, engine='openpyxl', engine_kwargs={'data_only': True})

  temp_filepath = filepath.with_name('temp_' + filepath.stem + filepath.suffix)

  df.to_excel(temp_filepath, index=False, engine='openpyxl')

  wb = load_workbook(temp_filepath, read_only=True, data_only=True)

  ws = wb[wb.sheetnames[0]]
  xl_lastrow = xl_get_last_row(ws)

  printLabel("✅ VOR Loaded")

  sections = []
  works_count = 0
  mat_count = 0
  
  
  for i, row in enumerate(ws.iter_rows(min_row=1, max_row=xl_lastrow, values_only=True)):
    if row[0] == "№ п.п.": 
      head_row = i + 1

      pp_col = get_column(row, '№ п.п.', 0)
      name_col = get_column(row, 'Наименование работ', 1)
      meas_col = get_column(row, 'Ед. изм.', 2)
      amount_col = get_column(row, 'Объем', 3)
      type_col = get_column(row, 'Тип позиции', 4)
      break

  for i, row in enumerate(ws.iter_rows(min_row=1, max_row=xl_lastrow, values_only=True)):
    if re.match("Раздел ", str(row[0])) : 
      sect_name = re.search(r"Раздел \d+. (.*)", row[0]).group(1)
      sections.append([sect_name, i + 1])
    else:
      if str(row[type_col]).lower() == 'материал':
        mat_count += 1
      elif str(row[type_col]).lower() == 'работа':
        works_count += 1
    
  if statistics_lbl:
    stat_text = (f'Найдено разделов: {len(sections) if len(sections) > 0 else 'Разделы не найдены'}\n')
    stat_text = stat_text + f'Всего работ: {works_count}\n'
    stat_text = stat_text + f'Всего материалов: {mat_count}'
    statistics_lbl.setText(stat_text)
    QApplication.processEvents()

  if file_data:
    file_data['file_path'] = temp_filepath
    file_data['sections'] = sections

  wb.close()

def create_query_xl(section_name=None, category_name=None, object_name=None):
  if section_name is None:
    return None
  
  
  printLabel("💹 Creating Excel query")

  folder_date = date_to_filename(datetime.now())

  xl_output_path = project_path / 'data' / 'xl_output' / object_name / str(folder_date + " " + section_name)

  if not os.path.exists(xl_output_path):
    os.makedirs(xl_output_path)

  xl_filename = xl_output_path / ('Запрос ' + category_name + f' от {datetime.now().strftime("%Y-%m-%d")}.xlsx')

  if os.path.isfile(xl_filename):
    print("⚠️ File already exists")
    query_vor_head_template(xl_filename)
    return xl_filename
  
  query_vor_head_template(xl_filename)
  
  print("✅ File created")
  return xl_filename

def insert_data_xl(ws, data, include_works=True):
  
  printLabel("Setting query data")
  ws_lastrow = ws.max_row
  row_increment = 1

  for val in data:
    item = val['data']
    if include_works == False and item['type'].lower() == 'работа':
      continue
    if item['include'] == False:
      continue

    for col, val in enumerate(['code', 'name', 'chars', 'measure', 'amount'], 1):
      # if col > 5:
      #   break
      cell = ws.cell(row=ws_lastrow + row_increment, column=col, value=item[val])
      if item['type'].lower() == 'работа':
        cell.font, cell.fill, cell.alignment, cell.border = xl_format_cell('work_cell')
      else:
        cell.font, cell.fill, cell.alignment, cell.border = xl_format_cell('material_cell')
      
    for col in range(6, 9):
      cell = ws.cell(row=ws_lastrow + row_increment, column=col)
      if col == 7: 
        cell.font, cell.fill, cell.alignment, cell.border = xl_format_cell('comment_cell')
      else:
        cell.font, cell.fill, cell.alignment, cell.border = xl_format_cell('price_cell')
        cell.number_format = style.numbers.FORMAT_NUMBER_COMMA_SEPARATED1
    row_increment += 1
    
    


  # for index, val in enumerate(data):
  #   ws_lastrow = ws.max_row
  #   for col, val_item in enumerate(val, 1):
  #     cell = ws.cell(row=ws_lastrow + 1, column=col, value=val_item)
  #     if index == 0:
  #       cell.font, cell.fill, cell.alignment, cell.border = xl_format_cell('work_cell')
  #     else:
  #       cell.font, cell.fill, cell.alignment, cell.border = xl_format_cell('material_cell')
  #   for col in range(len(val) + 1, ws.max_column + 1):
  #     cell = ws.cell(row=ws_lastrow + 1, column=col)
  #     if col == ws.max_column + 1:
  #       cell.font, cell.fill, cell.alignment, cell.border = xl_format_cell('comment_cell')
  #     else:
  #       cell.font, cell.fill, cell.alignment, cell.border = xl_format_cell('price_cell')
  #       cell.number_format = style.numbers.FORMAT_NUMBER_COMMA_SEPARATED1


def xl_format_cell(format_type):
  if format_type == 'header':
    cell_font = style.Font(bold=True, size=12, color=pyxl_color('#FFFFFF'), name='Arial')
    cell_fill = style.PatternFill(fgColor=pyxl_color("#100068"), fill_type='solid')
    cell_align = style.Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell_borders = style.Border(
      left=style.Side(style='thin'),
      right=style.Side(style='thin'),
      top=style.Side(style='thin'),
      bottom=style.Side(style='thick')
    )
  elif format_type == 'material_cell':
    cell_font = style.Font(size=9, color=pyxl_color("#000000"), name='Arial')
    cell_fill = style.PatternFill(fgColor=pyxl_color("#FFFFFF"), fill_type='solid')
    cell_align = style.Alignment(horizontal='right', vertical='center', wrap_text=True)
    cell_borders = style.Border(
        left=style.Side(style='thin'),
        right=style.Side(style='thin'),
        top=style.Side(style='thin'),
        bottom=style.Side(style='thin')
      )
  elif format_type == 'work_cell':
    cell_font = style.Font(italic=True, size=9, color=pyxl_color("#000000"), name='Arial')
    cell_fill = style.PatternFill(fgColor=pyxl_color("#FFFFFF"), fill_type='solid')
    cell_align = style.Alignment(horizontal='left', vertical='center', wrap_text=True)
    cell_borders = style.Border(
        left=style.Side(style='thin'),
        right=style.Side(style='thin'),
        top=style.Side(style='thin'),
        bottom=style.Side(style='thin')
      )
  elif format_type == 'comment_cell':
    cell_font = style.Font(italic=False, size=9, color=pyxl_color("#000000"), name='Arial')
    cell_fill = style.PatternFill(fgColor=pyxl_color("#B0FFAD"), fill_type='solid')
    cell_align = style.Alignment(horizontal='left', vertical='center', wrap_text=True)
    cell_borders = style.Border(
        left=style.Side(style='thin'),
        right=style.Side(style='thin'),
        top=style.Side(style='thin'),
        bottom=style.Side(style='thin')
      )
  elif format_type == 'price_cell':
    cell_font = style.Font(italic=False, size=9, color=pyxl_color("#000000"), name='Arial')
    cell_fill = style.PatternFill(fgColor=pyxl_color("#B0FFAD"), fill_type='solid')
    cell_align = style.Alignment(horizontal='left', vertical='center', wrap_text=True)
    cell_borders = style.Border(
        left=style.Side(style='thin'),
        right=style.Side(style='thin'),
        top=style.Side(style='thin'),
        bottom=style.Side(style='thin')
      )

  return cell_font, cell_fill, cell_align, cell_borders

def query_vor_head_template(file_path):

  wb = Workbook()
  ws = wb.active

  header_text = ['п/п', 'Наименование позции', 'Тех.характеристики', 'Ед.изм', 'Объем', 'Цена за единицу', 'Стоимость', 'Примечание']

  from openpyxl.utils import get_column_letter

  for col_ind, val in enumerate(header_text, 1):
    cell = ws.cell(row=1, column=col_ind, value=val)
    cell.font, cell.fill, cell.alignment, cell.border = xl_format_cell('header')

    ws.column_dimensions[get_column_letter(col_ind)].width = len(val) * 2

  wb.save(file_path)

def choose_file():
  from tkinter import filedialog
  import tkinter as tk
  root = tk.Tk()
  root.withdraw()

  file_path = filedialog.askopenfilename(
    title="Choose Excel file"
    # filetypes=[("Excel files", "*.xlsx", "*xlsm", "*xls"), ("All files", "*.*")]
  )

  return file_path

if __name__ == "__main__":
  vor_file_path = Path(choose_file()) # project_path / 'data' / 'ВОР пример.xlsx' 
  get_vor_file(vor_file_path)
  
  
  printLabel("✅ Finished")