from apps.crm.services.model_services.company_service import CompanyService
from apps.crm.services.model_services.connection_service import ConnectionService

from openpyxl import Workbook
from openpyxl.styles import Alignment, PatternFill, Border, Side, Font
from io import BytesIO

def format_cell(cell, val=None, font=None, fill=None, border=None, align=None):
  if font:
    cell.font = font
  if border:
    cell.border = border
  if align:
    cell.alignment = align
  if fill:
    cell.fill = fill
  if val:
    cell.value = val

class ExportService:
  @staticmethod
  def get_object_brend_list(object_id):
    brend_list = ConnectionService.get_brend_list(object_id)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Бренд-лист'

    worksheet.append([
      'Система',
      'Оборудование/Материалы',
      'Бренд-лист объекта'
    ])
    header_font = Font(bold=True, italic=True)
    section_font = Font(bold=True)

    thin = Side(style='thin', color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    total_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_center = Alignment(horizontal="left", vertical="center", wrap_text=True)

    section_fill = PatternFill('solid', fgColor='FF8EA9DB')
    split_fill = PatternFill('solid', fgColor='FFB4C6E7')

    for adr in ['A1', 'B1', 'C1']:
      format_cell(worksheet[adr], font=header_font, align=total_center)

    sections = []
    categories = []
    materials = []

    for item in brend_list:
      sections.append(item.section)
      categories.append((item.section, item.material.parent.name))
      materials.append((item.section, item.material.parent.name, item.material.name))


    sections = list(set(sections))
    sections = sorted(sections)
    categories = list(set(categories))
    materials = list(set(materials))


    brends_output = {}
    for item in brend_list:
      key = (
        item.section,
        item.material.parent.name,
        item.material.name
      )
      brends_output.setdefault(key, []).append(item.company.name)
      
    current_row = 2

    for section in sections:
      format_cell(worksheet[f"A{current_row}"], val=section, border=border, align=total_center, fill=section_fill, font=section_font)
      worksheet.merge_cells(f"A{current_row}:C{current_row}")
      
      current_row += 1
      
      section_categories = [
        category
        for category in categories
        if category[0] == section
      ]
      for idx, category in enumerate(section_categories):
        format_cell(worksheet[f"A{current_row}"], val=category[1], align=total_center, border=border)
        category_start = current_row
        category_materials = [material for material in materials if section==material[0] and category[1]==material[1]]
        for material in category_materials:
          format_cell(worksheet[f'B{current_row}'], val=material[2], align=left_center, border=border)
          material_brends = ', '.join(brends_output.get(material, []))
          format_cell(worksheet[f'C{current_row}'], val=material_brends, align=total_center, border=border)
          current_row += 1
        category_end = current_row - 1
        worksheet.merge_cells(f'A{category_start}:A{category_end}')
        

        # РАЗДЕЛИТЕЛЬ КАТЕГОРИЙ
        if idx+1 < len(section_categories):
          current_row = worksheet.max_row + 1
          format_cell(worksheet[f"A{current_row}"], border=border, fill=split_fill)
          worksheet.merge_cells(f"A{current_row}:C{current_row}")
        
        current_row = worksheet.max_row + 1
      current_row = worksheet.max_row + 1

    worksheet.column_dimensions['A'].width = 25
    worksheet.column_dimensions['B'].width = 60
    worksheet.column_dimensions['C'].width = 60

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    
    return output
  
  @staticmethod
  def get_companies_output(companies_ids):
    output_companies = CompanyService.companies_filter_ids(companies_ids).prefetch_related('contacts', 'materials')
    output_data = []

    all_connections = ConnectionService.get_companies_contacts(companies_ids)
    connections_by_company = {}

    for conn in all_connections:
      connections_by_company.setdefault(conn.company_id, []).append(conn)
  
    for company in output_companies:
      comp_cont = connections_by_company.get(company.id, [])
        
      contacts = company.contacts.all()
    
      materials = company.materials.all()
      materials_list = '; '.join([material.name for material in materials])
      company_contacts = {
        'name': [],
        'mail': [],
        'phone': []
      }
      if company.mail:
        company_contacts['mail'].append(company.mail)
      if company.phone:
        company_contacts['phone'].append(company.phone)

      for cont in contacts:
        connection = next((item for item in comp_cont if item.contact_id == cont.id), None)
        company_contacts['name'].append(f'{cont.last_name} {cont.first_name}')

        if connection and connection.mail: 
          company_contacts['mail'].append(connection.mail)
        if connection and connection.phone:
          company_contacts['phone'].append(connection.phone)

      output_data.append({
        'company_name': company.name,
        'names': '; '.join(company_contacts['name']),
        'mails': '; '.join(company_contacts['mail']),
        'phones': '; '.join(company_contacts['phone']),
        'materials': materials_list
      })

    return output_data